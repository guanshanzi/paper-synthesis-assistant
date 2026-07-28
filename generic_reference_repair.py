from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
import re
import sys
import traceback


BASE_DIR = Path(__file__).parent
OUT_DIR = BASE_DIR / "output"

INPUT_CANDIDATES = [
    OUT_DIR / "综述定稿候选_硬约束通过.docx",
    OUT_DIR / "综述最终稿_V10_引文顺序修正版.docx",
    OUT_DIR / "综述最终稿_V9_标点修正版.docx",
    OUT_DIR / "综述最终稿_V8.docx",
    OUT_DIR / "综述定稿候选_V7_小修定稿.docx",
    OUT_DIR / "综述定稿候选_V6_微修稿.docx",
    OUT_DIR / "综述定稿候选_V5_中修稿.docx",
    OUT_DIR / "综述定稿候选_V4_补齐一级标题.docx",
]

OUTPUT_DOCX = OUT_DIR / "综述最终稿_通用引文修复版.docx"
REPORT_PATH = OUT_DIR / "通用引文修复报告.txt"

TITLE_KEYWORDS = ["防止", "背景", "综述", "研究", "对策"]


def find_input_docx():
    for p in INPUT_CANDIDATES:
        if p.exists():
            return p

    docxs = sorted(OUT_DIR.glob("*.docx"), key=lambda x: x.stat().st_mtime, reverse=True)

    exclude_keywords = [
        "审稿",
        "评分",
        "材料包",
        "参考文献清单",
        "资料卡片",
        "精简",
    ]

    prefer_keywords = [
        "综述",
        "定稿",
        "最终稿",
        "候选",
    ]

    for p in docxs:
        if p.name.startswith("~$"):
            continue
        if any(k in p.name for k in exclude_keywords):
            continue
        if any(k in p.name for k in prefer_keywords):
            return p

    for p in docxs:
        if p.name.startswith("~$"):
            continue
        if any(k in p.name for k in exclude_keywords):
            continue
        return p

    return None


def read_texts(path):
    doc = Document(path)
    return [p.text.strip() for p in doc.paragraphs if p.text.strip()]


def split_body_and_refs(texts):
    body = []
    refs = []
    in_refs = False

    for t in texts:
        if t == "参考文献":
            in_refs = True
            continue

        if in_refs:
            refs.append(t)
        else:
            body.append(t)

    return body, refs


def extract_ref_map(refs):
    ref_map = {}
    for r in refs:
        m = re.match(r"^\[(\d+)\]\s*(.+)", r.strip())
        if m:
            ref_map[m.group(1)] = r.strip()
    return ref_map


def normalize_cite_content(content):
    content = content.replace("，", ",").replace("、", ",").replace("—", "-")
    content = re.sub(r"\s+", "", content)
    return content


def expand_cite_content(content):
    content = normalize_cite_content(content)
    nums = []

    if not content:
        return nums

    parts = [p for p in content.split(",") if p]

    for p in parts:
        if "-" in p:
            a, b = p.split("-", 1)
            if a.isdigit() and b.isdigit():
                start = int(a)
                end = int(b)
                for n in range(min(start, end), max(start, end) + 1):
                    nums.append(str(n))
        elif p.isdigit():
            nums.append(str(int(p)))

    return nums


def compress_nums(nums):
    nums = sorted(set(int(n) for n in nums))
    if not nums:
        return ""

    ranges = []
    start = nums[0]
    prev = nums[0]

    for n in nums[1:]:
        if n == prev + 1:
            prev = n
        else:
            if start == prev:
                ranges.append(str(start))
            else:
                ranges.append(f"{start}-{prev}")
            start = n
            prev = n

    if start == prev:
        ranges.append(str(start))
    else:
        ranges.append(f"{start}-{prev}")

    return ",".join(ranges)


def remove_invalid_cites_from_text(text, existing_ref_nums):
    """
    删除正文中参考文献表不存在的引文编号。
    如果一个引文组里只有无效编号，则删除整个引文括号。
    """
    removed = []

    def repl(m):
        raw = m.group(1)
        nums = expand_cite_content(raw)

        valid = []
        for n in nums:
            if n in existing_ref_nums:
                valid.append(n)
            else:
                removed.append(n)

        if not valid:
            return ""

        return "[" + compress_nums(valid) + "]"

    new_text = re.sub(r"\[([0-9,\-—，、\s]+)\]", repl, text)
    return new_text, removed


def build_renumber_map(body_text, ref_map):
    """
    删除无效引文后，按正文首次出现顺序重编号。
    """
    mapping = {}
    order_old = []
    next_num = 1

    for m in re.finditer(r"\[([0-9,\-—，、\s]+)\]", body_text):
        nums = expand_cite_content(m.group(1))

        for old in nums:
            if old not in ref_map:
                continue

            if old not in mapping:
                mapping[old] = str(next_num)
                order_old.append(old)
                next_num += 1

    return mapping, order_old


def replace_citations_by_mapping(text, mapping):
    def repl(m):
        old_nums = expand_cite_content(m.group(1))
        new_nums = []

        for old in old_nums:
            if old in mapping:
                new_nums.append(mapping[old])

        if not new_nums:
            return ""

        return "[" + compress_nums(new_nums) + "]"

    return re.sub(r"\[([0-9,\-—，、\s]+)\]", repl, text)


def renumber_refs(ref_map, mapping, order_old):
    new_refs = []

    for old in order_old:
        if old not in ref_map:
            continue

        new_num = mapping[old]
        old_ref = ref_map[old]
        new_ref = re.sub(r"^\[\d+\]", f"[{new_num}]", old_ref)
        new_refs.append(new_ref)

    return new_refs


def set_style(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5

    for style_name in ["Title", "Heading 1", "Heading 2", "标题 1", "标题 2"]:
        if style_name in doc.styles:
            s = doc.styles[style_name]
            s.font.name = "宋体"
            s._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            s.font.size = Pt(12)


def is_heading(text):
    return bool(re.match(r"^\d+(\.\d+)?\s+.+", text.strip()))


def write_docx(body_texts, refs):
    doc = Document()
    set_style(doc)

    for i, t in enumerate(body_texts):
        if i == 0:
            p = doc.add_heading(t, level=0)
        elif is_heading(t):
            level = 2 if re.match(r"^\d+\.\d+\s+", t) else 1
            p = doc.add_heading(t, level=level)
        else:
            p = doc.add_paragraph(t)

        for run in p.runs:
            run.font.name = "宋体"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            run.font.size = Pt(12)

    p = doc.add_heading("参考文献", level=1)
    for run in p.runs:
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(12)

    for r in refs:
        p = doc.add_paragraph(r)
        for run in p.runs:
            run.font.name = "宋体"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            run.font.size = Pt(12)

    doc.save(OUTPUT_DOCX)


def extract_cited_nums(text):
    cited = set()
    for m in re.finditer(r"\[([0-9,\-—，、\s]+)\]", text):
        for n in expand_cite_content(m.group(1)):
            cited.add(n)
    return cited


def write_report(input_path, removed_invalid, mapping, final_body_text, final_refs):
    ref_nums = set()
    for r in final_refs:
        m = re.match(r"^\[(\d+)\]", r)
        if m:
            ref_nums.add(m.group(1))

    cited_nums = extract_cited_nums(final_body_text)

    invalid_after = sorted([n for n in cited_nums if n not in ref_nums], key=lambda x: int(x))
    uncited_after = sorted([n for n in ref_nums if n not in cited_nums], key=lambda x: int(x))

    lines = []
    lines.append("通用引文修复报告")
    lines.append("=" * 50)
    lines.append(f"输入文件：{input_path}")
    lines.append(f"输出文件：{OUTPUT_DOCX}")
    lines.append("")
    lines.append("一、删除的无效正文引文编号")
    if removed_invalid:
        lines.append("、".join(sorted(set(removed_invalid), key=lambda x: int(x))))
    else:
        lines.append("无")
    lines.append("")
    lines.append("二、重编号映射")
    if mapping:
        for old, new in sorted(mapping.items(), key=lambda x: int(x[1])):
            lines.append(f"[{old}] -> [{new}]")
    else:
        lines.append("无")
    lines.append("")
    lines.append("三、修复后检查")
    lines.append(f"正文实际引用文献数量：{len(cited_nums)}")
    lines.append(f"参考文献数量：{len(ref_nums)}")
    lines.append(f"修复后非法引文：{', '.join(invalid_after) if invalid_after else '无'}")
    lines.append(f"修复后未引用参考文献：{', '.join(uncited_after) if uncited_after else '无'}")
    lines.append("")
    if not invalid_after and not uncited_after:
        lines.append("结论：通过。正文引文和参考文献表已重新对齐。")
    else:
        lines.append("结论：未通过。仍需人工或后续程序处理。")

    REPORT_PATH.write_text("\n".join(lines), encoding="utf-8")


def main():
    try:
        input_path = find_input_docx()
        if not input_path:
            print("没有找到可修复的 Word 文件。")
            sys.exit(1)

        print(f"正在读取：{input_path}")
        texts = read_texts(input_path)
        body_texts, refs = split_body_and_refs(texts)

        ref_map = extract_ref_map(refs)
        existing_ref_nums = set(ref_map.keys())

        if not ref_map:
            print("没有识别到参考文献表，无法修复。")
            sys.exit(1)

        removed_all = []
        cleaned_body_texts = []

        for t in body_texts:
            new_t, removed = remove_invalid_cites_from_text(t, existing_ref_nums)
            cleaned_body_texts.append(new_t)
            removed_all.extend(removed)

        cleaned_body_text = "\n".join(cleaned_body_texts)

        mapping, order_old = build_renumber_map(cleaned_body_text, ref_map)

        final_body_texts = [
            replace_citations_by_mapping(t, mapping)
            for t in cleaned_body_texts
        ]

        final_refs = renumber_refs(ref_map, mapping, order_old)
        final_body_text = "\n".join(final_body_texts)

        write_docx(final_body_texts, final_refs)
        write_report(input_path, removed_all, mapping, final_body_text, final_refs)

        print("通用引文修复完成。")
        print(f"Word：{OUTPUT_DOCX}")
        print(f"报告：{REPORT_PATH}")

    except Exception as e:
        err = f"程序出错：{e}\n\n{traceback.format_exc()}"
        print(err)
        REPORT_PATH.write_text(err, encoding="utf-8")
        sys.exit(1)


if __name__ == "__main__":
    main()