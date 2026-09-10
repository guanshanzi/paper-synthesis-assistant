from openai import OpenAI
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from pathlib import Path
import csv
import json
import re
import time
import traceback
import sys
import os


# =========================================================
# 一、API 与模型配置
# =========================================================

API_KEY = os.getenv("API_KEY") or os.getenv("OPENAI_API_KEY") or "请在界面填写API_KEY"
BASE_URL = os.getenv("BASE_URL") or os.getenv("OPENAI_BASE_URL") or "https://api.302.ai/v1"
MODEL = os.getenv("CARD_MODEL") or os.getenv("WRITE_MODEL") or "gpt-5.4"
SLEEP_SECONDS = 2

PAPER_TITLE = (os.getenv("PAPER_TITLE") or "").strip()
if not PAPER_TITLE:
    title_file = Path(__file__).parent / "论文题目.txt"
    if title_file.exists():
        PAPER_TITLE = title_file.read_text(encoding="utf-8", errors="replace").strip()
if not PAPER_TITLE:
    PAPER_TITLE = "未填写论文题目"


# =========================================================
# 二、文件夹设置
# =========================================================

BASE_DIR = Path(__file__).parent
CARDS_DIR = BASE_DIR / "cards_input"
OUT_DIR = BASE_DIR / "output"
OUT_DIR.mkdir(exist_ok=True)

MATRIX_CSV_PATH = OUT_DIR / "证据矩阵.csv"
COMPACT_WORD_PATH = OUT_DIR / "精简资料卡片汇总.docx"
PROGRESS_PATH = OUT_DIR / "matrix_progress.txt"
ERROR_LOG_PATH = OUT_DIR / "matrix_error_log.txt"

client = OpenAI(api_key=API_KEY, base_url=BASE_URL)


# =========================================================
# 三、Word 设置
# =========================================================

def set_word_style(doc):
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "宋体"
    normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal_style.font.size = Pt(12)
    normal_style.paragraph_format.line_spacing = 1.5


def create_or_load_compact_word():
    if COMPACT_WORD_PATH.exists():
        doc = Document(COMPACT_WORD_PATH)
        set_word_style(doc)
        return doc

    doc = Document()
    set_word_style(doc)
    doc.add_heading("精简资料卡片汇总", level=0)
    doc.add_paragraph("以下内容为对原始资料卡片进行二次压缩后的结果。")
    doc.save(COMPACT_WORD_PATH)
    return doc


def append_compact_card(doc, title, content):
    doc.add_page_break()
    doc.add_heading(title, level=1)

    for line in content.splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        run = p.add_run(line)
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(12)

    doc.save(COMPACT_WORD_PATH)


# =========================================================
# 四、进度和日志
# =========================================================

def load_progress():
    if not PROGRESS_PATH.exists():
        return set()
    return set(line.strip() for line in PROGRESS_PATH.read_text(encoding="utf-8").splitlines() if line.strip())


def mark_done(card_id):
    with PROGRESS_PATH.open("a", encoding="utf-8") as f:
        f.write(card_id + "\n")


def log_error(text):
    with ERROR_LOG_PATH.open("a", encoding="utf-8") as f:
        f.write(text + "\n\n")


# =========================================================
# 五、读取 Word 并拆分资料卡片
# =========================================================

def read_docx_cards(docx_path):
    doc = Document(docx_path)

    lines = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            lines.append(text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    text = p.text.strip()
                    if text:
                        lines.append(text)

    full_text = "\n".join(lines)

    def extract_title_from_card(part, idx):
        first_line = part.splitlines()[0].strip() if part.splitlines() else ""

        if re.search(r"资料卡片|文献卡片|论文资料卡片", first_line):
            return first_line[:120]

        title_patterns = [
            r"论文名称\s*[：:]\s*([^\n]+)",
            r"论文题名\s*[：:]\s*([^\n]+)",
            r"文献题名\s*[：:]\s*([^\n]+)",
            r"题名\s*[：:]\s*([^\n]+)",
            r"题目\s*[：:]\s*([^\n]+)",
        ]

        for pat in title_patterns:
            m = re.search(pat, part)
            if m:
                return m.group(1).strip()[:120]

        m = re.search(r"《([^》]{3,100})》", part)
        if m:
            return f"资料卡片{idx}：《{m.group(1)}》"

        if first_line:
            return f"资料卡片{idx}：{first_line[:80]}"

        return f"资料卡片{idx}"

    def make_cards_by_positions(positions, mode_name):
        cards = []
        if not positions:
            return cards

        positions = sorted(set(positions))

        for i, start in enumerate(positions):
            end = positions[i + 1] if i + 1 < len(positions) else len(full_text)
            part = full_text[start:end].strip()

            if len(part) < 200:
                continue

            cards.append({
                "title": extract_title_from_card(part, i + 1),
                "content": part,
                "mode": mode_name
            })

        return cards

    candidate_groups = []

    card_marker_pattern = (
        r"(?:资料卡片|文献资料卡片|论文资料卡片|文献卡片)"
        r"\s*[（(【\[]?\s*"
        r"(?:第\s*)?"
        r"[0-9０-９一二三四五六七八九十百千万]+"
        r"\s*[）)】\]]?"
        r"\s*[：:、.．]?"
    )
    positions_1 = [m.start() for m in re.finditer(card_marker_pattern, full_text)]
    candidate_groups.append(make_cards_by_positions(positions_1, "资料卡片编号"))

    basic_info_pattern = (
        r"(?:^|\n)"
        r"(?:一|1|Ⅰ|I)"
        r"\s*[、.．]\s*"
        r"(?:论文基本信息|文献基本信息|基本信息|论文信息)"
    )
    positions_2 = []
    for m in re.finditer(basic_info_pattern, full_text):
        pos = m.start()
        if full_text[pos:pos + 1] == "\n":
            pos += 1
        positions_2.append(pos)
    candidate_groups.append(make_cards_by_positions(positions_2, "论文基本信息"))

    title_pattern = (
        r"(?:^|\n)"
        r"(?:论文名称|论文题名|文献题名|题名|题目)"
        r"\s*[：:]"
    )
    positions_3 = []
    for m in re.finditer(title_pattern, full_text):
        pos = m.start()
        if full_text[pos:pos + 1] == "\n":
            pos += 1
        positions_3.append(pos)
    candidate_groups.append(make_cards_by_positions(positions_3, "论文题名"))

    pdf_score_pattern = r"(?:^|\n).{1,160}\.pdf\s*[｜|]\s*相关度评分\s*[：:]?\s*\d"
    positions_4 = []
    for m in re.finditer(pdf_score_pattern, full_text):
        pos = m.start()
        if full_text[pos:pos + 1] == "\n":
            pos += 1
        positions_4.append(pos)
    candidate_groups.append(make_cards_by_positions(positions_4, "PDF相关度标题"))

    candidate_groups = [g for g in candidate_groups if len(g) > 0]

    if not candidate_groups:
        print("没有识别出任何资料卡片。")
        return []

    cards = max(candidate_groups, key=lambda g: len(g))
    mode_name = cards[0].get("mode", "未知方式")

    for c in cards:
        c.pop("mode", None)

    max_card = max(cards, key=lambda x: len(x["content"]))
    avg_len = sum(len(c["content"]) for c in cards) // max(len(cards), 1)

    print(f"拆分方式：{mode_name}")
    print(f"拆分检查：共 {len(cards)} 张卡片。")
    print(f"最长卡片：{max_card['title'][:100]}")
    print(f"最长卡片字符数：{len(max_card['content'])}")
    print(f"平均卡片字符数：{avg_len}")

    if len(max_card["content"]) > 50000:
        raise RuntimeError(
            "拆分仍然异常：最长卡片超过 50000 字。程序已停止，避免继续消耗额度。"
        )

    return cards


# =========================================================
# 六、调用模型，把长卡片压成证据矩阵
# =========================================================

def call_model(prompt):
    response = client.chat.completions.create(
        model=MODEL,
        messages=[
            {"role": "user", "content": prompt}
        ]
    )
    content = response.choices[0].message.content
    return content.strip() if content else ""


def parse_json_safely(result, title):
    try:
        text = result.strip()
        text = re.sub(r"^```json", "", text).strip()
        text = re.sub(r"^```", "", text).strip()
        text = re.sub(r"```$", "", text).strip()

        start = text.find("{")
        end = text.rfind("}")
        if start != -1 and end != -1 and end > start:
            text = text[start:end + 1]

        return json.loads(text)

    except Exception:
        return {
            "title": title,
            "authors": "",
            "year": "",
            "source": "",
            "relevance_score": "",
            "priority": "解析失败",
            "chapter": "",
            "problem_supported": "",
            "countermeasure_supported": "",
            "core_view": result[:300],
            "usable_material": "",
            "limitation": "模型输出未能解析为JSON，需要人工查看。",
            "compact_card": result[:1000]
        }


def summarize_card_to_matrix(title, content):
    content = content[:12000]

    prompt = f"""
你正在帮助我整理已有的论文资料卡片。注意：这不是重新总结论文原文，而是把“已有资料卡片”压缩成可用于综述写作的证据矩阵。

我的论文主题是：《{PAPER_TITLE}》。

请根据下面这张资料卡片，提取其对我论文写作最有用的信息。

请严格输出 JSON，不要输出其他说明。字段如下：

{{
  "title": "文献题名或文件名",
  "authors": "作者，若资料卡片中没有则留空",
  "year": "年份，若没有则留空",
  "source": "期刊或来源，若没有则留空",
  "relevance_score": "0-5分",
  "priority": "核心/重要/一般/背景/不建议使用",
  "chapter": "适合放入论文的章节或部分，可多选；请根据当前论文题目与资料卡片内容概括，不要套用旧项目的固定分类",
  "problem_supported": "这篇文献能支撑的困境，不超过120字",
  "countermeasure_supported": "这篇文献能支撑的对策，不超过120字",
  "core_view": "最核心观点，不超过150字",
  "usable_material": "可直接用于论文的材料，不超过200字",
  "limitation": "使用局限，不超过120字",
  "compact_card": "压缩版资料卡片，300-500字"
}}

原始资料卡片标题：
{title}

原始资料卡片内容：
{content}
"""

    result = call_model(prompt)
    data = parse_json_safely(result, title)

    if not data.get("title"):
        data["title"] = title

    return data


# =========================================================
# 七、CSV 输出
# =========================================================

def init_matrix_csv():
    if MATRIX_CSV_PATH.exists():
        return

    with MATRIX_CSV_PATH.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow([
            "文献编号",
            "文献题名",
            "作者",
            "年份",
            "来源",
            "相关度评分",
            "使用优先级",
            "适合章节",
            "可支撑困境",
            "可支撑对策",
            "核心观点",
            "可用材料",
            "使用局限"
        ])


def get_next_ref_number():
    if not MATRIX_CSV_PATH.exists():
        return 1

    max_num = 0

    with MATRIX_CSV_PATH.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            ref = str(row.get("文献编号", "")).strip()
            m = re.search(r"\d+", ref)
            if m:
                max_num = max(max_num, int(m.group()))

    return max_num + 1


def append_matrix_csv(ref_num, data):
    with MATRIX_CSV_PATH.open("a", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow([
            ref_num,
            data.get("title", ""),
            data.get("authors", ""),
            data.get("year", ""),
            data.get("source", ""),
            data.get("relevance_score", ""),
            data.get("priority", ""),
            data.get("chapter", ""),
            data.get("problem_supported", ""),
            data.get("countermeasure_supported", ""),
            data.get("core_view", ""),
            data.get("usable_material", ""),
            data.get("limitation", "")
        ])


# =========================================================
# 八、主程序
# =========================================================

def main():
    if "在这里" in API_KEY or not API_KEY.strip():
        print("请先在代码顶部填写 API_KEY。")
        sys.exit(1)

    if not CARDS_DIR.exists():
        print(f"没有找到 cards_input 文件夹：{CARDS_DIR}")
        print("请新建 cards_input 文件夹，并把已有资料卡片 Word 放进去。")
        sys.exit(1)

    docx_files = sorted(CARDS_DIR.glob("*.docx"))

    if not docx_files:
        print("cards_input 文件夹里没有 docx 文件。")
        sys.exit(1)

    init_matrix_csv()
    compact_doc = create_or_load_compact_word()
    done = load_progress()

    print(f"共发现 {len(docx_files)} 个资料卡片 Word。")

    for docx_path in docx_files:
        print(f"\n正在读取：{docx_path.name}")
        cards = read_docx_cards(docx_path)
        print(f"拆出 {len(cards)} 张资料卡片。")

        for idx, card in enumerate(cards, start=1):
            title = card["title"]
            content = card["content"]
            card_id = f"{docx_path.name}::{title}"

            if card_id in done:
                print(f"跳过已处理：{title}")
                continue

            print(f"\n正在二次整理：{title}")

            try:
                ref_num = get_next_ref_number()
                data = summarize_card_to_matrix(title, content)

                append_matrix_csv(ref_num, data)

                compact_title = f"[{ref_num}] {data.get('title', title)}"
                compact_card = data.get("compact_card", "")
                append_compact_card(compact_doc, compact_title, compact_card)

                mark_done(card_id)

                print(f"已写入证据矩阵和精简卡片。文献编号：[{ref_num}]")
                time.sleep(SLEEP_SECONDS)

            except Exception as e:
                error_info = (
                    f"处理卡片时出错，程序已停止。\n"
                    f"Word：{docx_path.name}\n"
                    f"卡片：{title}\n"
                    f"错误信息：{str(e)}\n"
                    f"{traceback.format_exc()}"
                )

                print(error_info)
                log_error(error_info)

                append_compact_card(
                    compact_doc,
                    "程序停止说明",
                    f"处理到《{title}》时发生错误，程序已停止。\n\n错误信息：{str(e)}\n\n前面结果已保存。"
                )

                sys.exit(1)

    print("\n处理结束。")
    print(f"证据矩阵：{MATRIX_CSV_PATH}")
    print(f"精简卡片 Word：{COMPACT_WORD_PATH}")
    print(f"进度记录：{PROGRESS_PATH}")
    print(f"错误日志：{ERROR_LOG_PATH}")


if __name__ == "__main__":
    main()
